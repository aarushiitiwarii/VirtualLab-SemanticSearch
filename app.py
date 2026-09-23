import streamlit as st
import pandas as pd
import plotly.graph_objects as go
import re
import os
from io import BytesIO
import numpy as np

os.environ.setdefault("TRANSFORMERS_NO_TF", "1")
os.environ.setdefault("USE_TF", "0")

from sentence_transformers import SentenceTransformer
from sklearn.metrics.pairwise import cosine_similarity


# ---------------- PAGE SETUP ----------------

st.set_page_config(
    page_title="Dense Embedding-Based Semantic Search",
    page_icon="🔎",
    layout="wide"
)

# ---------------- LOAD CUSTOM CSS ----------------

with open("style.css") as f:
    st.markdown(
        f"<style>{f.read()}</style>",
        unsafe_allow_html=True
    )
    
# ---------------- MODEL ----------------

@st.cache_resource
def load_model():
    return SentenceTransformer("all-MiniLM-L6-v2")


# ---------------- SESSION STATE ----------------

if "documents" not in st.session_state:
    st.session_state.documents = []

if "document_embeddings" not in st.session_state:
    st.session_state.document_embeddings = None

if "search_results" not in st.session_state:
    st.session_state.search_results = None

if "last_query" not in st.session_state:
    st.session_state.last_query = ""

documents = st.session_state.documents


# ---------------- KEYWORD SEARCH ----------------

def keyword_search(query, documents, top_k):

    query_words = set(
        re.findall(r"\b\w+\b", query.lower())
    )

    scores = []

    for document in documents:

        document_words = set(
            re.findall(r"\b\w+\b", document.lower())
        )

        matching_words = query_words.intersection(
            document_words
        )

        scores.append(len(matching_words))

    ranked_indices = sorted(
        range(len(scores)),
        key=lambda i: scores[i],
        reverse=True
    )

    return ranked_indices[:top_k], scores


def extract_uploaded_text(uploaded_file):

    file_name = uploaded_file.name.lower()
    file_bytes = uploaded_file.getvalue()

    if file_name.endswith(".pdf"):

        try:
            from pypdf import PdfReader

            reader = PdfReader(BytesIO(file_bytes))
            return "\n".join(
                page.extract_text() or ""
                for page in reader.pages
            ).strip()

        except ImportError:
            raise ValueError(
                "PDF support requires the pypdf package. Install it with: pip install pypdf"
            )

    if file_name.endswith((".jpg", ".jpeg", ".png")):

        try:
            from PIL import Image
            pytesseract = __import__("pytesseract")

            image = Image.open(BytesIO(file_bytes))
            return pytesseract.image_to_string(image).strip()

        except ImportError:
            raise ValueError(
                "Image OCR requires Pillow and pytesseract. Install them with: "
                "pip install Pillow pytesseract"
            )

    if file_name.endswith(".docx"):

        try:
            from docx import Document

            document = Document(BytesIO(file_bytes))
            return "\n".join(
                paragraph.text
                for paragraph in document.paragraphs
            ).strip()

        except ImportError:
            raise ValueError(
                "DOCX support requires python-docx. Install it with: pip install python-docx"
            )

    return file_bytes.decode("utf-8", errors="ignore").strip()


# ---------------- TITLE ----------------

st.title("Dense Embedding-Based Semantic Search")

st.write("Virtual Laboratory Experiment")


# ---------------- SIDEBAR ----------------

st.sidebar.title("Experiment Sections")

section = st.sidebar.radio(
    "Select a section:",
    [
        "Purpose",
        "Theory",
        "Simulation",
        "Quiz",
        "Report Generation",
        "Certificate",
        "References"
    ]
)


st.header(section)


# ---------------- PURPOSE ----------------

if section == "Purpose":

    st.subheader("Aim")

    st.write(
        "To develop a semantic search system using dense vector embeddings "
        "to retrieve documents that are conceptually similar to a user's query, "
        "even when the query and document use different words."
    )

    st.subheader("Experiment Background")

    st.write(
        "Traditional keyword search looks for exact words shared by a query and a "
        "document. This approach is useful for precise terms, but it may fail when "
        "the user uses a synonym, changes the word order, or describes an idea in "
        "different language. Semantic search addresses this limitation by representing "
        "both queries and documents as dense numerical vectors. Their meanings can then "
        "be compared using a mathematical similarity measure."
    )

    st.write(
        "This virtual experiment demonstrates the complete retrieval process on a small "
        "collection of documents related to artificial intelligence and information "
        "retrieval. It also compares the embedding-based result with a simple keyword "
        "baseline, allowing the learner to observe when meaning-based retrieval produces "
        "a different ranking from exact word matching."
    )

    st.subheader("Objectives")

    st.markdown("""
    - Understand why exact keyword matching may miss conceptually related information.
    - Generate dense vector embeddings for the document collection and user query.
    - Interpret a fixed-length vector as a learned numerical representation of text.
    - Calculate cosine similarity between the query vector and document vectors.
    - Rank documents according to their similarity scores.
    - Retrieve the top-k documents most relevant to the query.
    - Compare semantic search with traditional keyword-based search.
    - Explain how the embedding model and document collection affect retrieval quality.
    """)

    st.subheader("Learning Outcomes")

    st.write(
        "After completing the experiment, the learner should be able to describe the "
        "role of an embedding model in information retrieval, write and explain the "
        "cosine-similarity formula, interpret a ranked search result, and distinguish "
        "semantic similarity from simple word overlap. The learner should also be able "
        "to identify why a high similarity score indicates relevance according to the "
        "model, but does not by itself guarantee that two texts are identical or correct."
    )

    st.subheader("Expected Outcome")

    st.write(
        "The simulation should return a ranked list of documents for the entered query. "
        "The highest-ranked document is the one with the greatest cosine similarity to "
        "the query embedding. The comparison table should help demonstrate that semantic "
        "search can find related content even when fewer exact query words appear in the "
        "document."
    )


# ---------------- THEORY ----------------

if section == "Theory":

    st.subheader("1. What is Semantic Search?")

    st.write(
        "Semantic search retrieves documents according to the meaning and context "
        "of the text rather than relying only on exact word matches. For example, "
        "a query about computers understanding language may retrieve a document "
        "about natural language processing even when the words in the query and "
        "document are not identical. This makes semantic search useful when users "
        "express the same idea in different ways."
    )

    st.write(
        "In this experiment, the searchable collection contains short documents. "
        "The same procedure can be applied to lecture notes, webpages, support "
        "articles, or scientific papers. Each document is converted into a vector "
        "before searching, so the search operation compares numerical representations "
        "of meaning instead of comparing raw strings alone."
    )

    st.subheader("2. Dense Embeddings")

    st.write(
        "A dense embedding represents a document or query as a fixed-length numerical "
        "vector. The model used here, all-MiniLM-L6-v2, maps each text into a "
        "384-dimensional vector. Every coordinate contributes some learned information "
        "about the text, such as its topic, context, or relationship to other words. "
        "The coordinates are not individual word counts; they are learned features."
    )

    st.latex(r"d_i = f(\text{document}_i), \qquad q = f(\text{query})")

    st.write(
        "Here, f represents the embedding model, d_i is the vector for document i, "
        "and q is the vector for the user's query. Because the same model encodes "
        "both documents and queries, their vectors can be compared in the same vector "
        "space. Documents that are close to the query in this space are treated as "
        "more relevant."
    )

    st.subheader("3. Similarity Measurement")

    st.write(
        "The similarity between the query vector and each document vector is measured "
        "using cosine similarity. Cosine similarity measures the angle between two "
        "vectors, so it focuses on their direction and is less affected by the "
        "absolute length of the vectors. A score close to 1 indicates strong alignment, "
        "a score near 0 indicates weak alignment, and a negative score indicates "
        "opposing directions."
    )

    st.latex(
        r"\operatorname{cosine\_similarity}(q,d_i) "
        r"= \frac{q \cdot d_i}{\|q\|\,\|d_i\|}"
    )

    st.write(
        "The numerator q · d_i is the dot product. If q = (q_1, q_2, ..., q_n) "
        "and d_i = (d_1, d_2, ..., d_n), the dot product is calculated by multiplying "
        "corresponding coordinates and adding the results. The denominator normalizes "
        "the result using the Euclidean length (L2 norm) of each vector."
    )

    st.latex(
        r"q \cdot d_i = \sum_{j=1}^{n} q_jd_{ij}, \qquad "
        r"\|q\| = \sqrt{\sum_{j=1}^{n}q_j^2}, \qquad "
        r"\|d_i\| = \sqrt{\sum_{j=1}^{n}d_{ij}^2}"
    )

    st.subheader("Formula Summary")

    st.write(
        "The complete calculation can also be understood as a sequence of four "
        "mathematical operations: encode the text, normalize the vectors, calculate "
        "their alignment, and rank the documents by the resulting score."
    )

    st.latex(
        r"\hat{q} = \frac{q}{\|q\|}, \qquad "
        r"\hat{d}_i = \frac{d_i}{\|d_i\|}"
    )

    st.write(
        "The hat symbol represents a normalized vector. After normalization, cosine "
        "similarity is equivalent to the dot product of the normalized query and "
        "document vectors."
    )

    st.latex(
        r"\operatorname{score}_i = \hat{q} \cdot \hat{d}_i "
        r"= \operatorname{cosine\_similarity}(q,d_i)"
    )

    st.latex(
        r"\operatorname{ranking} = "
        r"\operatorname{argsort}(\operatorname{score}_1, "
        r"\operatorname{score}_2, \ldots, \operatorname{score}_m) "
        r"\text{ in descending order}"
    )

    st.write(
        "If there are m documents, the system calculates m scores and returns the "
        "first k document indices from this descending ranking. This is the mathematical "
        "meaning of top-k retrieval used in the simulation."
    )

    st.subheader("4. Retrieval")

    st.write(
        "After calculating one score for every document, the system sorts the scores "
        "from highest to lowest. The top-k documents are returned, where k is selected "
        "by the user in the Simulation section. The value of k controls how many results "
        "are displayed, but it does not change the embedding or the similarity scores."
    )

    st.subheader("5. Complete Experimental Workflow")

    st.markdown("""
    1. Store the document collection.
    2. Encode every document once to create dense document embeddings.
    3. Encode the user's query with the same embedding model.
    4. Compute cosine similarity between the query vector and every document vector.
    5. Sort documents by decreasing similarity score.
    6. Display the top-k results and compare them with keyword matching.
    """)

    st.subheader("6. Semantic Search Compared with Keyword Search")

    st.write(
        "The keyword baseline counts the distinct query words that also occur in each "
        "document. Its simple score is therefore:"
    )

    st.latex(
        r"\operatorname{keyword\_score}(q,d_i) "
        r"= |\operatorname{words}(q) \cap \operatorname{words}(d_i)|"
    )

    st.write(
        "Keyword matching is transparent and can be effective when exact terminology "
        "matters, but it may miss synonyms and related concepts. Semantic search can "
        "retrieve conceptually related text, although its result depends on the model, "
        "the quality of the document collection, and the meaning captured by the embeddings."
    )

    st.subheader("7. Worked Calculation")

    st.write(
        "For one query and one document, the simulation substitutes the generated vector "
        "values into the cosine formula. It first calculates the dot product, then the "
        "two vector norms, and finally divides the numerator by the denominator. The "
        "result is the score used for ranking. The Simulation section shows these actual "
        "values for the documents selected in a search."
    )

    st.subheader("Basic Workflow")

    st.markdown("""
    **Documents → Dense Embeddings → Vector Representation**

    **User Query → Dense Embedding → Similarity Calculation → Ranking → Top-k Results**
    """)

    st.subheader("Key Terms")

    terms = {
        "Dense Embedding":
            "A numerical vector representing the meaning of text.",

        "Semantic Search":
            "Search based on meaning and context.",

        "Cosine Similarity":
            "A measure of similarity between two vectors.",

        "Top-k Retrieval":
            "Selecting the k documents with the highest similarity scores."
    }

    for term, definition in terms.items():

        st.markdown(
            f"**{term}:** {definition}"
        )

    st.subheader("8. Limitations and Interpretation")

    st.write(
        "A high cosine score should be interpreted as strong similarity according to the "
        "embedding model, not as proof that two documents are identical or factually "
        "equivalent. Short texts may provide limited context, and a model can reflect "
        "biases or gaps in the data used to train it. In a production system, results "
        "should be evaluated with representative queries and human relevance judgments."
    )


# ---------------- SIMULATION ----------------

if section == "Simulation":

    st.subheader("Semantic Search Simulation")


    # STEP 1
    st.subheader("Step 1: Document Collection")

    st.write(
        "Add one or more text documents to create the collection used in this experiment. "
        "Each document should contain enough information to represent a meaningful topic "
        "or idea."
    )

    new_document = st.text_area(
        "Enter a document:",
        placeholder="Example: Natural language processing helps computers understand human language.",
        height=110,
        key="new_document"
    )

    uploaded_files = st.file_uploader(
        "Or upload documents and images:",
        type=[
            "pdf",
            "jpg",
            "jpeg",
            "png",
            "txt",
            "md",
            "csv",
            "json",
            "html",
            "docx"
        ],
        accept_multiple_files=True,
        help="PDF and DOCX files use text extraction. JPG and PNG files use OCR when pytesseract is installed."
    )

    add_document, clear_documents = st.columns([1, 1])

    with add_document:

        if st.button("Add Document", use_container_width=True):

            if not new_document.strip():

                st.warning("Enter some text before adding a document.")

            else:

                st.session_state.documents.append(new_document.strip())
                st.session_state.document_embeddings = None
                st.session_state.search_results = None
                st.success(
                    f"Document D{len(st.session_state.documents)} added to the collection."
                )
                st.rerun()

    if uploaded_files and st.button(
        "Import Uploaded Files",
        use_container_width=True
    ):

        imported_count = 0
        import_errors = []

        for uploaded_file in uploaded_files:

            try:
                extracted_text = extract_uploaded_text(uploaded_file)

                if extracted_text:
                    st.session_state.documents.append(extracted_text)
                    imported_count += 1
                else:
                    import_errors.append(
                        f"{uploaded_file.name}: no readable text was found"
                    )

            except ValueError as error:
                import_errors.append(f"{uploaded_file.name}: {error}")

        if imported_count:
            st.session_state.document_embeddings = None
            st.session_state.search_results = None
            st.success(f"Imported {imported_count} file(s) into the collection.")

        for import_error in import_errors:
            st.warning(import_error)

        if imported_count:
            st.rerun()

    with clear_documents:

        if st.button("Clear Collection", use_container_width=True):

            st.session_state.documents = []
            st.session_state.document_embeddings = None
            st.session_state.search_results = None
            st.session_state.last_query = ""
            st.rerun()

    documents = st.session_state.documents

    if not documents:

        st.info(
            "Your collection is empty. Add documents above before generating embeddings."
        )

    for i, document in enumerate(documents, start=1):

        st.write(
            f"**D{i}:** {document}"
        )


    # STEP 2
    st.subheader("Step 2: Generate Dense Embeddings")

    st.write(
        "The embedding model converts every document into a 384-dimensional vector. "
        "The vectors are not displayed in full because that would require showing 384 "
        "numbers per document, but the calculation below uses every coordinate."
    )

    st.latex(
        r"d_i = f(\text{document}_i), \qquad "
        r"q = f(\text{query}), \qquad "
        r"\operatorname{score}(q,d_i) = "
        r"\frac{\sum_{j=1}^{384}q_jd_{ij}}{"
        r"\sqrt{\sum_{j=1}^{384}q_j^2}\;"
        r"\sqrt{\sum_{j=1}^{384}d_{ij}^2}}"
    )

    if st.button("Generate Document Embeddings", disabled=not documents):

        with st.spinner("Loading the embedding model and generating vectors..."):
            model = load_model()
            st.session_state.document_embeddings = model.encode(documents)

        st.success(
            "Dense embeddings generated successfully!"
        )

        st.write(
            f"Each document has been converted into a "
            f"{st.session_state.document_embeddings.shape[1]}-dimensional vector."
        )


    # STEP 3
    st.subheader("Step 3: Enter Search Query")

    query = st.text_input(
        "Enter your search query:",
        placeholder="Example: How do computers understand human language?"
    )

    max_results = max(1, min(5, len(documents)))

    top_k = st.slider(
        "Number of results to retrieve:",
        min_value=1,
        max_value=max_results,
        value=min(3, max_results)
    )


    # STEP 4
    st.subheader("Step 4: Retrieve Similar Documents")

    if st.button("Search"):

        if not query.strip():

            st.warning(
                "Please enter a search query."
            )

        elif not documents:

            st.warning(
                "Please add at least one document before performing a search."
            )

        elif st.session_state.document_embeddings is None:

            st.warning(
                "Please generate the document embeddings "
                "before performing a search."
            )

        else:

            # Generate query embedding
            with st.spinner("Encoding the query and calculating similarities..."):
                model = load_model()
                query_embedding = model.encode([query])


            # Calculate cosine similarity
            similarity_scores = cosine_similarity(
                query_embedding,
                st.session_state.document_embeddings
            )[0]


            # Rank documents
            ranked_indices = similarity_scores.argsort()[::-1]


            # Select top-k documents
            top_indices = ranked_indices[:top_k]


            st.success(
                "Semantic search completed!"
            )


            # ---------------- SEMANTIC RESULTS ----------------

            st.subheader("Search Results")

            results = []

            for rank, index in enumerate(
                top_indices,
                start=1
            ):

                results.append({
                    "Rank": rank,
                    "Document": f"D{index + 1}",
                    "Similarity Score": round(
                        float(similarity_scores[index]),
                        4
                    )
                })


            results_df = pd.DataFrame(
                results
            )
            st.session_state.search_results = results_df
            st.session_state.last_query = query

            st.dataframe(
                results_df,
                use_container_width=True,
                hide_index=True
            )

            # ---------------- FORMULA TRACE ----------------

            st.subheader("Calculation Trace")

            st.write(
                "The following worked calculations use the actual query and document "
                "embeddings generated in this search. The first five vector components "
                "are shown as a readable sample, while the dot product, norms, and final "
                "score are calculated using all 384 components."
            )

            st.latex(
                r"\operatorname{cosine\_similarity}(q,d_i) "
                r"= \frac{q \cdot d_i}{\|q\|\,\|d_i\|}"
            )

            query_vector = query_embedding[0]

            for rank, index in enumerate(top_indices, start=1):

                document_vector = st.session_state.document_embeddings[index]
                dot_product = float(np.dot(query_vector, document_vector))
                query_norm = float(np.linalg.norm(query_vector))
                document_norm = float(np.linalg.norm(document_vector))
                calculated_score = dot_product / (query_norm * document_norm)

                with st.expander(f"Rank {rank}: D{index + 1} calculation"):

                    st.write(f"**Document:** {documents[index]}")
                    st.write(
                        f"Query vector sample (first 5 of 384): "
                        f"`{np.round(query_vector[:5], 4).tolist()}`"
                    )
                    st.write(
                        f"Document vector sample (first 5 of 384): "
                        f"`{np.round(document_vector[:5], 4).tolist()}`"
                    )

                    st.latex(
                        rf"q \cdot d_{{{index + 1}}} "
                        rf"= {dot_product:.6f}"
                    )
                    st.latex(
                        rf"\|q\| = {query_norm:.6f}, \qquad "
                        rf"\|d_{{{index + 1}}}\| = {document_norm:.6f}"
                    )
                    st.latex(
                        rf"\operatorname{{score}}(q,d_{{{index + 1}}}) "
                        rf"= \frac{{{dot_product:.6f}}}"
                        rf"{{{query_norm:.6f} \times {document_norm:.6f}}} "
                        rf"= {calculated_score:.6f}"
                    )

                    st.caption(
                        f"Displayed result score: {similarity_scores[index]:.6f}. "
                        f"The two values agree up to floating-point rounding."
                    )


            # ---------------- BAR CHART ----------------

            fig = go.Figure(
                go.Bar(
                    x=results_df["Document"],
                    y=results_df["Similarity Score"],
                    text=results_df["Similarity Score"],
                    textposition="auto"
                )
            )


            fig.update_layout(
                title="Semantic Similarity Scores",
                xaxis_title="Document",
                yaxis_title="Cosine Similarity",
                yaxis_range=[0, 1]
            )


            st.plotly_chart(
                fig,
                use_container_width=True
            )


            # ---------------- KEYWORD SEARCH ----------------

            st.subheader(
                "Keyword Search Comparison"
            )


            keyword_indices, keyword_scores = keyword_search(
                query,
                documents,
                top_k
            )


            for rank, index in enumerate(
                keyword_indices,
                start=1
            ):

                st.write(
                    f"**{rank}. D{index + 1}** - "
                    f"{documents[index]}"
                )

                st.write(
                    f"Keyword Matches: "
                    f"{keyword_scores[index]}"
                )

                st.divider()


            # ---------------- COMPARISON TABLE ----------------

            st.subheader(
                "Semantic vs Keyword Search"
            )


            comparison = []


            for rank in range(top_k):

                semantic_index = top_indices[rank]

                keyword_index = keyword_indices[rank]


                comparison.append({

                    "Rank":
                        rank + 1,

                    "Semantic Search":
                        f"D{semantic_index + 1}",

                    "Semantic Score":
                        round(
                            float(
                                similarity_scores[
                                    semantic_index
                                ]
                            ),
                            4
                        ),

                    "Keyword Search":
                        f"D{keyword_index + 1}",

                    "Keyword Matches":
                        keyword_scores[
                            keyword_index
                        ]
                })


            comparison_df = pd.DataFrame(
                comparison
            )


            st.dataframe(
                comparison_df,
                use_container_width=True,
                hide_index=True
            )
            st.subheader("Observation")

            best_semantic_index = top_indices[0]

            best_keyword_index = keyword_indices[0]

            st.info(
                f"Semantic search ranked D{best_semantic_index + 1} as the "
                f"most relevant document with a cosine similarity score of "
                f"{similarity_scores[best_semantic_index]:.4f}."
            )

            if best_semantic_index != best_keyword_index:

                st.write(
                    "The top result differs between semantic search and keyword "
                    "search. This shows that semantic search can identify "
                    "conceptually related documents even when the exact query "
                    "words are not present."
                )

            else:

                st.write(
                    "Both methods returned the same top document for this query. "
                    "However, semantic search ranks documents using their "
                    "meaning represented by dense embeddings."
                )


# ---------------- QUIZ ----------------

if section == "Quiz":

    st.subheader("Semantic Search Quiz")

    quiz_questions = [
        {
            "question": "What is the main purpose of semantic search?",
            "options": [
                "To match only exact keywords",
                "To retrieve documents based on meaning",
                "To store documents in a database",
                "To convert documents into images"
            ],
            "answer": "To retrieve documents based on meaning"
        },
        {
            "question": "What is a dense embedding?",
            "options": [
                "A database table",
                "A numerical vector representing text",
                "A keyword list",
                "A text file"
            ],
            "answer": "A numerical vector representing text"
        },
        {
            "question": "Which model is used in this experiment to generate embeddings?",
            "options": [
                "all-MiniLM-L6-v2",
                "ResNet-50",
                "BERT-Base-Cased",
                "GPT-2"
            ],
            "answer": "all-MiniLM-L6-v2"
        },
        {
            "question": "What does cosine similarity measure?",
            "options": [
                "The length of a document",
                "The number of words in a query",
                "The similarity between two vectors",
                "The number of documents"
            ],
            "answer": "The similarity between two vectors"
        },
        {
            "question": "What does a higher cosine similarity score generally indicate?",
            "options": [
                "Greater semantic similarity",
                "More spelling errors",
                "A longer document",
                "Fewer keywords"
            ],
            "answer": "Greater semantic similarity"
        },
        {
            "question": "What is the purpose of Top-k retrieval?",
            "options": [
                "To retrieve the k most relevant documents",
                "To remove all documents",
                "To generate k embeddings",
                "To count keywords"
            ],
            "answer": "To retrieve the k most relevant documents"
        },
        {
            "question": "What is the main difference between keyword search and semantic search?",
            "options": [
                "Keyword search uses images",
                "Semantic search uses exact words only",
                "Keyword search focuses on word matching while semantic search considers meaning",
                "There is no difference"
            ],
            "answer": "Keyword search focuses on word matching while semantic search considers meaning"
        },
        {
            "question": "What is generated for the user's search query?",
            "options": [
                "A database",
                "A dense embedding",
                "A PDF file",
                "A new document collection"
            ],
            "answer": "A dense embedding"
        },
        {
            "question": "How are documents ranked in semantic search?",
            "options": [
                "By document length",
                "By alphabetical order",
                "By similarity score",
                "By upload time"
            ],
            "answer": "By similarity score"
        },
        {
            "question": "Why can semantic search find conceptually related documents?",
            "options": [
                "Because embeddings capture semantic information",
                "Because every document contains the same words",
                "Because documents are sorted alphabetically",
                "Because it ignores the query"
            ],
            "answer": "Because embeddings capture semantic information"
        }
    ]


    # ---------------- QUESTIONS ----------------

    for i, question in enumerate(quiz_questions, start=1):

        with st.container(border=True):

            st.markdown(
                f'<div class="quiz-question-label">Question {i:02d}</div>',
                unsafe_allow_html=True
            )

            st.markdown(
                f'<div class="quiz-question-text">{question["question"]}</div>',
                unsafe_allow_html=True
            )

            st.radio(
                "Select your answer:",
                question["options"],
                index=None,
                key=f"quiz_{i}"
            )


    # ---------------- SUBMIT ----------------

    if st.button("Submit Quiz"):

        score = 0

        for i, question in enumerate(quiz_questions, start=1):

            selected_answer = st.session_state.get(
                f"quiz_{i}"
            )

            if selected_answer == question["answer"]:

                score += 1


        st.session_state.quiz_score = score

        st.success(
            f"Quiz submitted! Your score is {score}/10.",
            icon=None
        )


        # ---------------- FEEDBACK ----------------

        st.subheader("Quiz Feedback")

        for i, question in enumerate(quiz_questions, start=1):

            selected_answer = st.session_state.get(
                f"quiz_{i}"
            )

            if selected_answer == question["answer"]:

                st.markdown(
                    f"""
                    <div class="feedback-row feedback-correct">
                        <span class="feedback-number">{i:02d}</span>
                        <span class="feedback-question">Question {i}</span>
                        <span class="feedback-status">Correct</span>
                    </div>
                    """,
                    unsafe_allow_html=True
                )

            else:

                st.markdown(
                    f"""
                    <div class="feedback-row feedback-review">
                        <span class="feedback-number">{i:02d}</span>
                        <span class="feedback-question">Question {i}</span>
                        <span class="feedback-status">Review</span>
                    </div>
                    <div class="feedback-answer">
                        Correct answer: <strong>{question['answer']}</strong>
                    </div>
                    """,
                    unsafe_allow_html=True
                )

# ---------------- REPORT GENERATION ----------------

if section == "Report Generation":

    st.subheader("Experiment Report")

    student_name = st.text_input(
        "Student Name"
    )

    roll_number = st.text_input(
        "Roll Number"
    )

    st.write("### Experiment Details")

    st.write(
        "**Experiment:** Dense Embedding-Based Semantic Search"
    )

    st.write(
        "**Aim:** To develop a semantic search system using dense vector "
        "embeddings to retrieve conceptually similar documents."
    )

    st.subheader("Search Details")

    if st.session_state.search_results is not None:

        st.write(
            f"**Search Query:** {st.session_state.last_query}"
        )

        st.write(
            f"**Documents in Collection:** {len(documents)}"
        )

        st.write(
            f"**Results Retrieved:** "
            f"{len(st.session_state.search_results)}"
        )

        st.write("### Retrieved Documents")

        st.dataframe(
            st.session_state.search_results,
            use_container_width=True,
            hide_index=True
        )

    else:

        st.info(
            "Complete a semantic search in the Simulation section "
            "to include search results in the report."
        )
    st.subheader("Quiz Performance")

    if "quiz_score" in st.session_state:

        st.write(
            f"**Quiz Score:** "
            f"{st.session_state.quiz_score}/10"
        )

    else:

        st.info(
            "Complete the quiz to include the quiz score in the report."
        )

    st.subheader("Generate Report")

    if st.button("Generate PDF Report"):

        if not student_name.strip() or not roll_number.strip():

            st.warning(
                "Please enter your name and roll number first."
            )

        elif st.session_state.search_results is None:

            st.warning(
                "Please complete the semantic search before generating the report."
            )

        else:

            from fpdf import FPDF
            from datetime import datetime

            pdf = FPDF()
            pdf.add_page()

            pdf.set_font("Arial", "B", 16)
            pdf.cell(
                0,
                10,
                "Virtual Laboratory Experiment Report",
                ln=True,
                align="C"
            )

            pdf.ln(8)

            pdf.set_font("Arial", "B", 12)
            pdf.cell(
                0,
                8,
                "Dense Embedding-Based Semantic Search",
                ln=True
            )

            pdf.ln(5)

            pdf.set_font("Arial", "", 11)

            pdf.cell(
                0,
                8,
                f"Student Name: {student_name}",
                ln=True
            )

            pdf.cell(
                0,
                8,
                f"Roll Number: {roll_number}",
                ln=True
            )

            pdf.cell(
                0,
                8,
                f"Date: {datetime.now().strftime('%d-%m-%Y')}",
                ln=True
            )

            pdf.ln(5)

            pdf.set_font("Arial", "B", 12)
            pdf.cell(
                0,
                8,
                "Experiment Details",
                ln=True
            )

            pdf.set_font("Arial", "", 11)

            pdf.multi_cell(
                0,
                7,
                "Aim: To develop a semantic search system using "
                "dense vector embeddings to retrieve conceptually "
                "similar documents."
            )

            pdf.ln(4)

            pdf.cell(
                0,
                8,
                f"Search Query: {st.session_state.last_query}",
                ln=True
            )

            pdf.cell(
                0,
                8,
                f"Documents in Collection: {len(documents)}",
                ln=True
            )

            pdf.ln(5)

            pdf.set_font("Arial", "B", 12)
            pdf.cell(
                0,
                8,
                "Retrieved Documents",
                ln=True
            )

            pdf.set_font("Arial", "", 11)

            for _, row in st.session_state.search_results.iterrows():

                pdf.multi_cell(
                    180,
                    7,
                    f"Rank {row['Rank']} - "
                    f"{row['Document']} - "
                    f"Similarity Score: {row['Similarity Score']}"
                )

                pdf.ln(2)

            pdf.set_font("Arial", "B", 12)
            pdf.cell(
                0,
                8,
                "Quiz Performance",
                ln=True
            )

            pdf.set_font("Arial", "", 11)

            quiz_score = st.session_state.get(
                "quiz_score",
                "Not completed"
            )

            pdf.cell(
                0,
                8,
                f"Quiz Score: {quiz_score}/10",
                ln=True
            )

            pdf.ln(5)

            pdf.set_font("Arial", "B", 12)
            pdf.cell(
                0,
                8,
                "Conclusion",
                ln=True
            )

            pdf.set_font("Arial", "", 11)

            pdf.multi_cell(
                0,
                7,
                "The experiment demonstrates how dense embeddings "
                "can be used to represent text numerically and retrieve "
                "documents based on semantic similarity."
            )

            pdf_bytes = bytes(pdf.output())

            st.download_button(
                label="Download PDF Report",
                data=pdf_bytes,
                file_name="semantic_search_report.pdf",
                mime="application/pdf"
            )

# ---------------- CERTIFICATE ----------------

if section == "Certificate":

    st.subheader("Virtual Laboratory Certificate")

    certificate_name = st.text_input(
        "Enter your name for the certificate:"
    )

    if st.button("Generate Certificate"):

        if not certificate_name.strip():

            st.warning(
                "Please enter your name first."
            )

        else:

            from fpdf import FPDF

            pdf = FPDF(
                orientation="L",
                unit="mm",
                format="A4"
            )
            pdf.set_margins(18, 16, 18)
            pdf.add_page()

            page_width = 297
            page_height = 210

            pdf.set_draw_color(35, 55, 70)
            pdf.set_line_width(1.2)
            pdf.rect(10, 10, page_width - 20, page_height - 20)

            pdf.set_draw_color(173, 123, 53)
            pdf.set_line_width(0.5)
            pdf.rect(15, 15, page_width - 30, page_height - 30)

            pdf.set_text_color(35, 55, 70)
            pdf.set_font("Arial", "B", 13)
            pdf.cell(0, 12, "VIRTUAL LABORATORY", align="C", ln=True)

            pdf.set_text_color(173, 123, 53)
            pdf.set_font("Times", "B", 28)
            pdf.cell(0, 18, "CERTIFICATE OF COMPLETION", align="C", ln=True)

            pdf.set_draw_color(173, 123, 53)
            pdf.line(106, 58, 191, 58)
            pdf.ln(10)

            pdf.set_text_color(78, 85, 90)
            pdf.set_font("Arial", "", 12)
            pdf.cell(0, 9, "This certificate is proudly presented to", align="C", ln=True)
            pdf.ln(3)

            pdf.set_text_color(35, 55, 70)
            pdf.set_font("Times", "B", 27)
            pdf.cell(0, 18, certificate_name.strip(), align="C", ln=True)

            pdf.set_draw_color(173, 123, 53)
            pdf.line(82, 100, 215, 100)
            pdf.ln(9)

            pdf.set_text_color(78, 85, 90)
            pdf.set_font("Arial", "", 12)
            pdf.cell(0, 8, "for successfully completing the virtual laboratory experiment", align="C", ln=True)

            pdf.set_text_color(35, 55, 70)
            pdf.set_font("Arial", "B", 16)
            pdf.cell(0, 12, "Dense Embedding-Based Semantic Search", align="C", ln=True)

            pdf.set_text_color(78, 85, 90)
            pdf.set_font("Arial", "", 11)
            pdf.multi_cell(
                0,
                7,
                "The learner demonstrated an understanding of dense embeddings, "
                "cosine similarity, semantic search, and top-k retrieval.",
                align="C"
            )

            pdf.set_y(166)
            pdf.set_draw_color(78, 85, 90)
            pdf.line(48, 178, 105, 178)
            pdf.line(192, 178, 249, 178)
            pdf.set_font("Arial", "", 9)
            pdf.set_text_color(78, 85, 90)
            pdf.set_xy(48, 180)
            pdf.cell(57, 6, "Virtual Laboratory", align="C")
            pdf.set_xy(192, 180)
            pdf.cell(57, 6, "Instructor", align="C")

            certificate_pdf = bytes(pdf.output())

            st.success(
                f"Certificate generated successfully for {certificate_name}!",
                icon=None
            )

            st.markdown(
                f"""
                ### Certificate of Completion

                This is to certify that

                ## **{certificate_name}**

                has successfully completed the Virtual Laboratory experiment

                **Dense Embedding-Based Semantic Search**

                and demonstrated an understanding of dense embeddings,
                cosine similarity, semantic search, and top-k retrieval.
                """
            )

            st.download_button(
                label="Download Certificate PDF",
                data=certificate_pdf,
                file_name="virtual_lab_certificate.pdf",
                mime="application/pdf"
            )


# ---------------- REFERENCES ----------------

if section == "References":

    st.subheader("References")

    references = [
        (
            "IIT Kharagpur Virtual Lab",
            "https://vlab.co.in/"
        ),
        (
            "Sentence Transformers Documentation",
            "https://sbert.net/"
        ),
        (
            "Scikit-learn Documentation",
            "https://scikit-learn.org/"
        ),
        (
            "Streamlit Documentation",
            "https://docs.streamlit.io/"
        ),
        (
            "FPDF2 Documentation",
            "https://py-pdf.github.io/fpdf2/"
        )
    ]

    for i, (name, link) in enumerate(references, start=1):

        st.markdown(
            f"{i}. **{name}**  \n"
            f"[{link}]({link})"
        )